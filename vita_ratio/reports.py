"""Экспорт дня/недели/месяца (PNG, DOCX, XLSX) и заметки с ответом ИИ-коуча.

Модуль намеренно самодостаточный: читает SQLite напрямую (как bot.py), без REST-слоя.
Формат отчёта: сводка КБЖУ/сон/план/вес + разбор коуча + заметки периода.
"""
from __future__ import annotations

import io
import sqlite3
from datetime import timedelta

from . import calc

ACCENT = "#B7F75C"
PERIODS = {"day": "День", "week": "Неделя", "month": "Месяц"}
MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля",
          "августа", "сентября", "октября", "ноября", "декабря"]


# --- данные -------------------------------------------------------------------

def _bounds(period: str, today: str) -> tuple[str, str, str]:
    day = calc.parse_date(today)
    if period == "week":
        start = day - timedelta(days=day.weekday())
        end = start + timedelta(days=6)
        return start.isoformat(), end.isoformat(), f"Неделя {day.isocalendar()[1]} ({start:%d.%m}–{end:%d.%m})"
    if period == "month":
        start = day.replace(day=1)
        end = (start + timedelta(days=32)).replace(day=1) - timedelta(days=1)
        return start.isoformat(), end.isoformat(), f"{MONTHS[start.month - 1].capitalize()} {start.year}"
    return day.isoformat(), day.isoformat(), f"День ({day:%d.%m.%Y})"


def collect(conn: sqlite3.Connection, user_id: int, period: str = "day", today: str | None = None) -> dict:
    """Сводка за период: профиль, цели, дни, сон, вес, план, занятия, заметки."""
    today = today or calc.local_today()
    start, end, label = _bounds(period, today)
    uid = int(user_id)
    profile = conn.execute("SELECT * FROM profiles WHERE user_id = ?", (uid,)).fetchone()
    targets = None
    if profile:
        try:
            targets = calc.targets(
                profile["sex"], profile["age"], profile["height_cm"], profile["weight_kg"],
                profile["activity"], profile["goal_rate"], profile["protein_per_kg"],
            ).to_dict()
        except Exception:
            targets = None

    rows = conn.execute(
        "SELECT date, COALESCE(SUM(kcal),0) kcal, COALESCE(SUM(protein),0) protein, "
        "COALESCE(SUM(fat),0) fat, COALESCE(SUM(carbs),0) carbs, COUNT(*) meals "
        "FROM meals WHERE user_id=? AND date BETWEEN ? AND ? GROUP BY date ORDER BY date",
        (uid, start, end),
    ).fetchall()
    day_totals = {str(r["date"]): dict(r) for r in rows}
    sleep = {str(r["date"]): float(r["duration_h"]) for r in conn.execute(
        "SELECT date, duration_h FROM sleep_logs WHERE user_id=? AND date BETWEEN ? AND ?", (uid, start, end))}
    weights = [dict(r) for r in conn.execute(
        "SELECT date, kg FROM weights WHERE user_id=? AND date BETWEEN ? AND ? ORDER BY date", (uid, start, end))]
    plans = conn.execute(
        "SELECT COUNT(*) total, COALESCE(SUM(done),0) done FROM day_plans "
        "WHERE user_id=? AND date BETWEEN ? AND ?", (uid, start, end)).fetchone()
    classes = conn.execute(
        "SELECT COUNT(*) FROM class_done WHERE user_id=? AND date BETWEEN ? AND ?", (uid, start, end)).fetchone()[0]
    notes = [dict(r) for r in conn.execute(
        "SELECT date, text, ai_text FROM notes WHERE user_id=? AND date BETWEEN ? AND ? ORDER BY created_at",
        (uid, start, end))]

    days = []
    cursor = calc.parse_date(start)
    while cursor <= calc.parse_date(end):
        key = cursor.isoformat()
        t = day_totals.get(key) or {}
        days.append({
            "date": key,
            "kcal": int(t.get("kcal") or 0),
            "protein": round(float(t.get("protein") or 0), 1),
            "fat": round(float(t.get("fat") or 0), 1),
            "carbs": round(float(t.get("carbs") or 0), 1),
            "meals": int(t.get("meals") or 0),
            "sleep_h": sleep.get(key),
        })
        cursor += timedelta(days=1)

    watched_days = [d for d in days if d["meals"] > 0]
    data = {
        "period": period,
        "label": label,
        "start": start,
        "end": end,
        "profile": dict(profile) if profile else None,
        "targets": targets,
        "days": days,
        "totals": {
            "kcal": sum(d["kcal"] for d in days),
            "protein": round(sum(d["protein"] for d in days), 1),
            "fat": round(sum(d["fat"] for d in days), 1),
            "carbs": round(sum(d["carbs"] for d in days), 1),
            "meals": sum(d["meals"] for d in days),
            "days_with_food": len(watched_days),
            "avg_kcal": int(round(sum(d["kcal"] for d in watched_days) / len(watched_days))) if watched_days else 0,
            "avg_sleep_h": round(sum(sleep.values()) / len(sleep), 1) if sleep else None,
        },
        "plans": {"total": int(plans["total"] or 0), "done": int(plans["done"] or 0)},
        "classes_done": int(classes or 0),
        "weight": {
            "start": round(float(weights[0]["kg"]), 1) if weights else None,
            "end": round(float(weights[-1]["kg"]), 1) if weights else None,
            "rows": weights,
        },
        "notes": notes,
    }
    data["comment"] = ai_comment(data)
    return data


def ai_comment(data: dict) -> str:
    """Локальный «ИИ-коуч»: короткий разбор цифр + конкретный совет (без внешнего API)."""
    t = data["totals"]
    targets = data.get("targets") or {}
    profile = data.get("profile") or {}
    tips: list[str] = []
    if t["days_with_food"] == 0:
        return ("Пока нет записей о еде за этот период. Начни с одного приёма пищи в день — "
                "и я покажу динамику калорий и белка.")

    avg = t["avg_kcal"]
    goal_kcal = int(targets.get("kcal") or 0)
    if goal_kcal:
        ratio = avg / goal_kcal
        if ratio < 0.8:
            tips.append(f"Среднесуточный калораж {avg} ккал — это {int(ratio * 100)}% от цели "
                        f"({goal_kcal}). Для набора стоит добавить ~{goal_kcal - avg} ккал: "
                        "ложка орехов, порция риса или протеиновый коктейль.")
        elif ratio > 1.2:
            tips.append(f"Среднесуточный калораж {avg} ккал — это {int(ratio * 100)}% от цели. "
                        "Профицит большой: следи, чтобы прибавка шла не только в жир — добавь шаги и тренировки.")
        else:
            tips.append(f"Калории в норме: в среднем {avg} ккал при цели {goal_kcal}. Так держать.")
    if targets.get("protein"):
        goal_protein = int(targets["protein"]) * max(1, t["days_with_food"])
        if t["protein"] < goal_protein * 0.8:
            tips.append(f"Белка {round(t['protein'])} г за период против ~{goal_protein} г — стоит "
                        "добавить творог, яйца или мясо в каждый приём.")
        else:
            tips.append(f"Белок почти по плану: {round(t['protein'])} г. Хорошо для роста мышц.")
    if t["avg_sleep_h"] is not None:
        if t["avg_sleep_h"] < 7:
            tips.append(f"Сон в среднем {t['avg_sleep_h']} ч — недосып снижает восстановление и аппетит; "
                        "постарайся убрать 30–40 минут из вечера и лечь раньше.")
        else:
            tips.append(f"Сон {t['avg_sleep_h']} ч в среднем — достаточно для восстановления.")
    if data["plans"]["total"]:
        pct = int(round(data["plans"]["done"] / data["plans"]["total"] * 100))
        if pct < 60:
            tips.append(f"План дня выполнен на {pct}%: выбери 1–2 главных блока и закрывай их первыми.")
        else:
            tips.append(f"План дня выполнен на {pct}% — дисциплина есть, продолжай.")
    if data["weight"]["start"] is not None and data["weight"]["end"] is not None:
        delta = round(data["weight"]["end"] - data["weight"]["start"], 1)
        goal_rate = float(profile.get("goal_rate") or 0.35)
        tips.append(f"Вес: {data['weight']['start']} → {data['weight']['end']} кг "
                    f"({delta:+}) при темпе цели {goal_rate} кг/нед.")
    note_text = (data["notes"][-1]["text"] if data["notes"] else "") or ""
    low = note_text.lower()
    if any(k in low for k in ("тяжело", "плохо", "устал", "болит", "стресс", "грустно")):
        tips.append("Из заметки видно, что непросто. Не жертвуй сном ради калорий: восстановление сейчас важнее плана — "
                    "сделай разгрузочный день и вернись в график завтра.")
    elif note_text:
        tips.append("Заметка учтена: держи фокус на том, что сработало в этом периоде.")
    return " ".join(tips[:5])


# --- заметки ------------------------------------------------------------------

def save_note(conn: sqlite3.Connection, user_id: int, text: str, today: str | None = None) -> str:
    """Сохраняет заметку за сегодня и возвращает ответ ИИ-коуча (он тоже пишется в БД)."""
    day = today or calc.local_today()
    data = collect(conn, int(user_id), "day", day)
    data["notes"] = data["notes"] + [{"date": day, "text": text, "ai_text": ""}]
    answer = ai_comment(data)
    conn.execute(
        "INSERT INTO notes (user_id, date, text, ai_text) VALUES (?, ?, ?, ?)",
        (int(user_id), day, text[:2000], answer),
    )
    conn.commit()
    return answer


def last_notes(conn: sqlite3.Connection, user_id: int, limit: int = 5) -> list[dict]:
    return [dict(r) for r in conn.execute(
        "SELECT date, text, ai_text FROM notes WHERE user_id=? ORDER BY id DESC LIMIT ?",
        (int(user_id), int(limit)))]


# --- форматы ------------------------------------------------------------------

def build(conn: sqlite3.Connection, user_id: int, period: str = "day", fmt: str = "png") -> tuple[str, bytes]:
    """Возвращает (имя файла, содержимое) для отправки в Telegram."""
    period = (period or "day").lower()
    fmt = (fmt or "png").lower()
    if period not in PERIODS:
        raise ValueError(f"period должен быть один из {sorted(PERIODS)}")
    if fmt not in ("png", "docx", "xlsx"):
        raise ValueError("fmt должен быть png, docx или xlsx")
    data = collect(conn, user_id, period)
    payload = {"png": _png, "docx": _docx, "xlsx": _xlsx}[fmt](data)
    return f"vita-ratio-{period}-{data['start']}.{fmt}", payload


def _rows(data: dict) -> list[tuple[str, str]]:
    t = data["totals"]
    targets = data.get("targets") or {}
    return [
        ("Период", data["label"]),
        ("Калории", f"{t['kcal']} ккал" + (f"  (цель {targets['kcal']})" if targets.get("kcal") else "")),
        ("Среднее в день", f"{t['avg_kcal']} ккал" + (" (по дням с едой)" if t["days_with_food"] else "")),
        ("Белок / жиры / углеводы", f"{t['protein']} / {t['fat']} / {t['carbs']} г"),
        ("Приёмов пищи", str(t["meals"])),
        ("Сон в среднем", f"{t['avg_sleep_h']} ч" if t["avg_sleep_h"] else "нет данных"),
        ("План дня", f"{data['plans']['done']}/{data['plans']['total']}"),
        ("Занятия отмечены", str(data["classes_done"])),
        ("Вес", f"{data['weight']['start']} → {data['weight']['end']} кг"
         if data["weight"]["start"] is not None else "нет данных"),
    ]


def _docx(data: dict) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Vita Ratio — отчёт", 0)
    doc.add_paragraph(data["label"])
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for key, value in _rows(data):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = key, value
    doc.add_heading("ИИ-разбор", level=1)
    doc.add_paragraph(data["comment"])
    if data["period"] != "day":
        doc.add_heading("По дням", level=1)
        table2 = doc.add_table(rows=1, cols=4)
        table2.style = "Light Grid Accent 1"
        for cell, title in zip(table2.rows[0].cells, ("Дата", "Ккал", "Белок, г", "Сон, ч")):
            cell.text = title
        for d in data["days"]:
            cells = table2.add_row().cells
            cells[0].text = d["date"]
            cells[1].text = str(d["kcal"])
            cells[2].text = str(d["protein"])
            cells[3].text = str(d["sleep_h"] or "—")
    if data["notes"]:
        doc.add_heading("Заметки", level=1)
        for note in data["notes"]:
            doc.add_paragraph(f"{note['date']}: {note['text']}")
            if note.get("ai_text"):
                doc.add_paragraph(f"ИИ: {note['ai_text']}")
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _xlsx(data: dict) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    summary = wb.active
    summary.title = "Сводка"
    summary.append(["Показатель", "Значение"])
    for key, value in _rows(data):
        summary.append([key, value])
    summary.append(["ИИ-разбор", data["comment"]])
    days_sheet = wb.create_sheet("Дни")
    days_sheet.append(["Дата", "Ккал", "Белок", "Жиры", "Углеводы", "Приёмов", "Сон, ч"])
    for d in data["days"]:
        days_sheet.append([d["date"], d["kcal"], d["protein"], d["fat"], d["carbs"], d["meals"], d["sleep_h"] or ""])
    notes_sheet = wb.create_sheet("Заметки")
    notes_sheet.append(["Дата", "Текст", "Ответ ИИ"])
    for note in data["notes"]:
        notes_sheet.append([note["date"], note["text"], note.get("ai_text", "")])
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _png(data: dict) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    def font(size: int):
        for path in ("C:/Windows/Fonts/arial.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"):
            try:
                return ImageFont.truetype(path, size)
            except OSError:
                continue
        return ImageFont.load_default()

    title_font, body_font, small_font = font(40), font(26), font(20)
    width = 1000
    rows = _rows(data)[1:]
    wrap_width = width - 120
    words, current, comment_lines = data["comment"].split(), "", []
    for word in words:
        probe = (current + " " + word).strip()
        if body_font.getlength(probe) > wrap_width:
            comment_lines.append(current)
            current = word
        else:
            current = probe
    if current:
        comment_lines.append(current)

    height = 200 + len(rows) * 46 + 60 + len(comment_lines) * 34 + 60
    image = Image.new("RGB", (width, height), "#0E1116")
    draw = ImageDraw.Draw(image)
    draw.rectangle([0, 0, width, 8], fill=ACCENT)
    y = 40
    draw.text((60, y), "Vita Ratio", font=title_font, fill="#FFFFFF")
    y += 52
    draw.text((60, y), data["label"], font=body_font, fill=ACCENT)
    y += 56
    for key, value in rows:
        draw.text((60, y), key, font=body_font, fill="#9AA4B2")
        draw.text((520, y), value, font=body_font, fill="#FFFFFF")
        y += 46
    y += 20
    draw.text((60, y), "ИИ-разбор", font=body_font, fill=ACCENT)
    y += 40
    for line in comment_lines:
        draw.text((60, y), line, font=small_font, fill="#E5E7EB")
        y += 34
    y += 20
    if data["notes"]:
        draw.text((60, y), "Заметки:", font=small_font, fill="#9AA4B2")
        y += 32
        for note in data["notes"][:3]:
            draw.text((60, y), f"• {note['date']}: {note['text'][:80]}", font=small_font, fill="#E5E7EB")
            y += 30
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()
